import os
from django.core.management.base import BaseCommand
from school_portal.models import Class, Student, Subject, Period, Mark
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime


class Command(BaseCommand):
    help = 'Génère un document statistique sur les résultats pour chaque classe'

    def add_arguments(self, parser):
        parser.add_argument(
            '--period',
            type=int,
            help='ID de la période pour filtrer les résultats'
        )
        parser.add_argument(
            '--output',
            type=str,
            default=None,
            help='Chemin de sortie du fichier (par défaut: Bureau)'
        )

    def handle(self, *args, **options):
        period_id = options.get('period')
        output_path = options.get('output')

        # Récupérer la période si spécifiée
        period = None
        if period_id:
            try:
                period = Period.objects.get(id=period_id)
                self.stdout.write(self.style.SUCCESS(f'Période sélectionnée: {period}'))
            except Period.DoesNotExist:
                self.stdout.write(self.style.ERROR(f'Période avec ID {period_id} non trouvée'))
                return
        else:
            # Utiliser la période active par défaut
            period = Period.objects.filter(is_active=True).first()
            if period:
                self.stdout.write(self.style.SUCCESS(f'Utilisation de la période active: {period}'))

        # Déterminer le chemin de sortie
        if not output_path:
            desktop_path = os.path.join(os.path.expanduser('~'), 'Desktop')
            if not os.path.exists(desktop_path):
                desktop_path = os.path.join(os.path.expanduser('~'), 'Bureau')

            timestamp = datetime.now().strftime('%Y%m%d_%H%M%S')
            period_suffix = f"_{period.name}" if period else ""
            output_path = os.path.join(desktop_path, f'Statistiques_Classes{period_suffix}_{timestamp}.xlsx')

        # Créer le classeur Excel
        wb = Workbook()
        wb.remove(wb.active)  # Supprimer la feuille par défaut

        # Récupérer toutes les classes
        classes = Class.objects.all().order_by('level', 'name')

        if not classes.exists():
            self.stdout.write(self.style.WARNING('Aucune classe trouvée'))
            return

        self.stdout.write(f'Génération des statistiques pour {classes.count()} classe(s)...')

        # Créer une feuille de synthèse
        summary_sheet = wb.create_sheet('Synthèse Générale')
        self._create_summary_sheet(summary_sheet, classes, period)

        # Créer une feuille par classe
        for classe in classes:
            self.stdout.write(f'  - Traitement de la classe {classe.name}...')
            sheet = wb.create_sheet(classe.name[:31])  # Excel limite à 31 caractères
            self._create_class_sheet(sheet, classe, period)

        # Sauvegarder le fichier
        try:
            wb.save(output_path)
            self.stdout.write(self.style.SUCCESS(f'\nFichier genere avec succes: {output_path}'))
        except Exception as e:
            self.stdout.write(self.style.ERROR(f'Erreur lors de la sauvegarde: {str(e)}'))

    def _create_summary_sheet(self, sheet, classes, period):
        """Crée la feuille de synthèse générale"""
        # Styles
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=12)
        title_font = Font(bold=True, size=14)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Titre
        period_text = f" - {period}" if period else ""
        sheet['A1'] = f'STATISTIQUES GÉNÉRALES PAR CLASSE{period_text}'
        sheet['A1'].font = title_font
        sheet['A1'].alignment = Alignment(horizontal='center')
        sheet.merge_cells('A1:H1')

        sheet['A2'] = f'Généré le: {datetime.now().strftime("%d/%m/%Y à %H:%M")}'
        sheet['A2'].alignment = Alignment(horizontal='center')
        sheet.merge_cells('A2:H2')

        # En-têtes
        headers = ['Classe', 'Niveau', 'Effectif', 'Moyenne Classe', 'Moyenne Min', 'Moyenne Max',
                   'Taux Réussite', 'Enseignant']
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=4, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center')
            cell.border = border

        # Données
        row = 5
        for classe in classes:
            students = classe.students.all()
            effectif = students.count()

            # Calculer les moyennes
            averages = []
            for student in students:
                avg = student.get_average(period=period)
                if avg is not None:
                    averages.append(avg)

            moyenne_classe = round(sum(averages) / len(averages), 2) if averages else 0
            moyenne_min = round(min(averages), 2) if averages else 0
            moyenne_max = round(max(averages), 2) if averages else 0
            taux_reussite = round((len([a for a in averages if a >= 10]) / len(averages) * 100), 2) if averages else 0

            teacher_name = f"{classe.teacher.name} {classe.teacher.first_name}" if classe.teacher else "Non assigné"

            data = [
                classe.name,
                classe.level,
                effectif,
                moyenne_classe,
                moyenne_min,
                moyenne_max,
                f"{taux_reussite}%",
                teacher_name
            ]

            for col, value in enumerate(data, start=1):
                cell = sheet.cell(row=row, column=col)
                cell.value = value
                cell.border = border
                cell.alignment = Alignment(horizontal='center' if col <= 7 else 'left', vertical='center')

            row += 1

        # Ajuster les largeurs de colonnes
        column_widths = [20, 15, 12, 15, 15, 15, 15, 25]
        for col, width in enumerate(column_widths, start=1):
            sheet.column_dimensions[get_column_letter(col)].width = width

    def _create_class_sheet(self, sheet, classe, period):
        """Crée une feuille détaillée pour une classe"""
        # Styles
        header_fill = PatternFill(start_color='366092', end_color='366092', fill_type='solid')
        subheader_fill = PatternFill(start_color='B8CCE4', end_color='B8CCE4', fill_type='solid')
        header_font = Font(bold=True, color='FFFFFF', size=11)
        title_font = Font(bold=True, size=13)
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Titre
        period_text = f" - {period}" if period else ""
        sheet['A1'] = f'STATISTIQUES DÉTAILLÉES - CLASSE {classe.name}{period_text}'
        sheet['A1'].font = title_font
        sheet['A1'].alignment = Alignment(horizontal='center')
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=6)

        # Informations de la classe
        row = 3
        info_data = [
            ('Niveau:', classe.level),
            ('Enseignant:', f"{classe.teacher.name} {classe.teacher.first_name}" if classe.teacher else "Non assigné"),
            ('Effectif:', classe.students.count()),
        ]

        for label, value in info_data:
            sheet[f'A{row}'] = label
            sheet[f'A{row}'].font = Font(bold=True)
            sheet[f'B{row}'] = value
            row += 1

        row += 1

        # Récupérer les matières de la classe
        subjects = Subject.objects.filter(classes=classe).order_by('name')

        # En-têtes du tableau
        headers = ['Matricule', 'Nom Complet', 'Moyenne Générale', 'Rang']

        # Ajouter les matières
        for subject in subjects:
            headers.append(subject.name)

        headers.append('Observation')

        # Écrire les en-têtes
        for col, header in enumerate(headers, start=1):
            cell = sheet.cell(row=row, column=col)
            cell.value = header
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
            cell.border = border

        row += 1

        # Récupérer les étudiants avec leurs moyennes
        students_data = []
        for student in classe.students.all():
            avg = student.get_average(period=period)
            if avg is not None:
                students_data.append({
                    'student': student,
                    'average': avg,
                    'rank': student.get_rank_in_class(period=period)
                })

        # Trier par moyenne décroissante
        students_data.sort(key=lambda x: x['average'], reverse=True)

        # Remplir les données des étudiants
        for item in students_data:
            student = item['student']
            moyenne_generale = item['average']
            rang = item['rank']

            # Informations de base
            data = [
                student.matricule,
                f"{student.name} {student.first_name}",
                moyenne_generale,
                rang if rang else "-"
            ]

            # Ajouter les moyennes par matière
            for subject in subjects:
                subject_avg = student.get_subject_average(subject, period=period)
                data.append(subject_avg if subject_avg is not None else "-")

            # Observation
            if moyenne_generale >= 14:
                observation = "Très Bien"
            elif moyenne_generale >= 12:
                observation = "Bien"
            elif moyenne_generale >= 10:
                observation = "Assez Bien"
            else:
                observation = "Insuffisant"

            data.append(observation)

            # Écrire la ligne
            for col, value in enumerate(data, start=1):
                cell = sheet.cell(row=row, column=col)
                cell.value = value
                cell.border = border

                # Centrage pour certaines colonnes
                if col in [1, 3, 4] or col > 4 and col < len(data):
                    cell.alignment = Alignment(horizontal='center', vertical='center')
                else:
                    cell.alignment = Alignment(horizontal='left', vertical='center')

                # Couleur pour la moyenne générale
                if col == 3:
                    if isinstance(value, (int, float)):
                        if value >= 14:
                            cell.fill = PatternFill(start_color='C6EFCE', end_color='C6EFCE', fill_type='solid')
                        elif value >= 10:
                            cell.fill = PatternFill(start_color='FFEB9C', end_color='FFEB9C', fill_type='solid')
                        else:
                            cell.fill = PatternFill(start_color='FFC7CE', end_color='FFC7CE', fill_type='solid')

            row += 1

        # Statistiques de la classe
        row += 1
        sheet[f'A{row}'] = 'STATISTIQUES DE LA CLASSE'
        sheet[f'A{row}'].font = Font(bold=True, size=11)
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=3)

        row += 1

        # Calculer les statistiques
        all_averages = [item['average'] for item in students_data]

        if all_averages:
            stats_data = [
                ('Moyenne de la classe:', round(sum(all_averages) / len(all_averages), 2)),
                ('Moyenne la plus élevée:', round(max(all_averages), 2)),
                ('Moyenne la plus basse:', round(min(all_averages), 2)),
                ('Nombre d\'élèves:', len(all_averages)),
                ('Nombre d\'élèves admis (≥10):', len([a for a in all_averages if a >= 10])),
                ('Taux de réussite:', f"{round(len([a for a in all_averages if a >= 10]) / len(all_averages) * 100, 2)}%"),
            ]

            for label, value in stats_data:
                sheet[f'A{row}'] = label
                sheet[f'A{row}'].font = Font(bold=True)
                sheet[f'B{row}'] = value
                row += 1

        # Ajuster les largeurs de colonnes
        sheet.column_dimensions['A'].width = 15
        sheet.column_dimensions['B'].width = 25
        sheet.column_dimensions['C'].width = 15
        sheet.column_dimensions['D'].width = 10

        # Largeur pour les matières
        for col in range(5, 5 + len(subjects)):
            sheet.column_dimensions[get_column_letter(col)].width = 12

        # Largeur pour observation
        sheet.column_dimensions[get_column_letter(5 + len(subjects))].width = 15
